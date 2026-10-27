"""Reading a source file: its text, exactly as the file holds it, and what it is.

The server reads the file itself, so the text the model labels from is the
file's own, not the assistant's copy of it. The archive keeps the file plus a
fingerprint of the extracted text (extractor name and sha1), not the text
itself. Text files, Markdown, HTML, PDF and Word (.docx). A PDF with no text
layer (a scan) has nothing to read: that is an error, never an empty source.
"""
from __future__ import annotations

import hashlib
import re
from html.parser import HTMLParser
from pathlib import Path

from hydra_mcp.core import IngestError

TYPES = {".txt": "text/plain", ".text": "text/plain", ".md": "text/markdown", ".markdown": "text/markdown",
         ".html": "text/html", ".htm": "text/html", ".pdf": "application/pdf",
         ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


class _Text(HTMLParser):
    SKIP, BLOCK = {"script", "style", "head", "noscript"}, {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4",
                                                              "h5", "h6", "section", "article", "blockquote", "td"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        self.skip += tag in self.SKIP
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        self.skip -= tag in self.SKIP and self.skip > 0
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def _decode(raw: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1"), "latin-1"


def _pdf(path: Path) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestError(f"{path.name} is password-protected; ask the person for an unlocked copy")
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
    except PdfReadError as e:
        raise IngestError(f"{path.name} could not be read as a PDF: {e}") from None
    return "\n\n".join(p for p in pages if p), len(pages)


def _docx(path: Path) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError
    try:
        d = docx.Document(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as e:
        raise IngestError(f"{path.name} could not be read as a Word document: {e}") from None
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def _tidy(text: str) -> str:
    text = re.sub(r"[ \t\r\f\v]+\n", "\n", text.replace("\r\n", "\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def read(path: Path) -> dict:
    """{name, type, size, sha1, extractor, text_sha1, pages, text}: the file, the tool
    that extracted its text, and the sha1 of that extracted text."""
    suffix = path.suffix.lower()
    if suffix not in TYPES:
        raise IngestError(f"{path.name}: hydra reads {', '.join(sorted(TYPES))} files"
                          + ("; save a .doc as .docx first" if suffix == ".doc" else ""))
    raw = path.read_bytes()
    pages = None
    if suffix == ".pdf":
        import pypdf
        text, pages = _pdf(path)
        extractor = f"pypdf {pypdf.__version__}"
    elif suffix == ".docx":
        import docx
        text = _docx(path)
        extractor = f"python-docx {docx.__version__}"
    elif suffix in (".html", ".htm"):
        p = _Text()
        p.feed(_decode(raw)[0])
        text = "".join(p.out)
        extractor = "html.parser"
    else:
        text, enc = _decode(raw)
        extractor = f"text read ({enc})"
    text = _tidy(text)
    if not text:
        raise IngestError(f"{path.name} has no text to read" + (
            ": it looks like a scan (images of pages, no text layer); it needs OCR first" if suffix == ".pdf" else ""))
    return {"name": path.name, "type": TYPES[suffix], "size": len(raw), "sha1": hashlib.sha1(raw).hexdigest(),
            "extractor": extractor, "text_sha1": hashlib.sha1(text.encode("utf-8")).hexdigest(),
            "pages": pages, "text": text}

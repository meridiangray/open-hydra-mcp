"""A file: read by the server, saved as the evidence with its original, cited by
every entity it names (new or already in the archive), and never saved twice."""
import hashlib
import json

import pytest

from hydra import layout
from hydra_mcp import core, server

TEXT = "Dana Hale, administrator of the Port Authority, said her deputy Rob Kerr would lead the inquiry."
RECORDS = [{"schema": "PublicBody", "properties": {"name": "Port Authority"}},
           {"schema": "Person", "properties": {"name": "Dana Hale"}},
           {"schema": "Person", "properties": {"name": "Rob Kerr"}},
           {"schema": "Employment", "properties": {"employee": "Dana Hale", "employer": "Port Authority", "role": "administrator"}},
           {"schema": "Employment", "properties": {"employee": "Rob Kerr", "employer": "Port Authority", "role": "deputy"}}]


def pdf(text: str | None) -> bytes:
    """A one-page PDF, with `text` on it, or with no text at all (like a scan)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)


@pytest.fixture
def archive(tmp_path):
    out = core.create_archive(str(tmp_path / "case"), name="case", confirm=True)
    return core.open_archive(out["archive"])


def claims(root):
    return [json.loads(l) for l in (root / layout.CLAIMS_JSONL).read_text().splitlines()]


def nodes_list(root):
    return list(json.loads((root / layout.READS["declared"]).read_text())["nodes"].values())


@pytest.mark.parametrize("name,content", [
    ("note.txt", TEXT.encode()),
    ("note.md", f"# Port inquiry\n\n{TEXT}\n".encode()),
    ("page.html", f"<html><head><style>p{{}}</style></head><body><h1>Port inquiry</h1><p>{TEXT}</p>"
                  "<script>var x = 1;</script></body></html>".encode()),
    ("filing.pdf", pdf(TEXT)),
])
def test_each_kind_of_file_is_read_by_the_server(archive, tmp_path, name, content):
    (tmp_path / name).write_bytes(content)
    got = server.read_file(str(tmp_path / name), archive=str(archive))
    assert TEXT in got["text"] and "var x" not in got["text"] and got["sha1"] == hashlib.sha1(content).hexdigest()
    assert got["already_in_archive"] is None and got["next_offset"] is None and "source=" in got["next"]


def test_a_word_document_is_read_with_its_tables(archive, tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph(TEXT)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Rob Kerr", "deputy"
    d.save(tmp_path / "memo.docx")
    got = server.read_file(str(tmp_path / "memo.docx"), archive=str(archive))
    assert TEXT in got["text"] and "Rob Kerr | deputy" in got["text"] and got["type"].endswith("document")


def test_unreadable_files_say_why(archive, tmp_path):
    (tmp_path / "scan.pdf").write_bytes(pdf(None))
    with pytest.raises(ValueError, match="looks like a scan"):
        server.read_file(str(tmp_path / "scan.pdf"), archive=str(archive))
    (tmp_path / "old.doc").write_bytes(b"x")
    with pytest.raises(ValueError, match="save a .doc as .docx"):
        server.read_file(str(tmp_path / "old.doc"), archive=str(archive))
    with pytest.raises(ValueError, match="not found"):
        server.read_file(str(tmp_path / "missing.pdf"), archive=str(archive))


def test_a_long_file_is_read_in_parts(archive, tmp_path):
    long = "\n\n".join(f"Paragraph {i}. " + " ".join(["word"] * 50) for i in range(400))
    (tmp_path / "long.txt").write_text(long)
    first = server.read_file(str(tmp_path / "long.txt"), archive=str(archive))
    rest = server.read_file(str(tmp_path / "long.txt"), offset=first["next_offset"], archive=str(archive))
    assert len(first["text"]) == core.READ_LIMIT and first["text"] + rest["text"] == long and rest["next_offset"] is None


def test_a_file_is_saved_as_the_evidence_with_its_original_and_exact_text(archive, tmp_path):
    content = pdf(TEXT)
    f = tmp_path / "filing.pdf"
    f.write_bytes(content)
    out = server.save(RECORDS, {"file": str(f), "publisher": "Acme Times"}, "filing", archive=str(archive))
    assert out["saved"] and out["files"] == ["collections/filing.json", ".hydra/sources/entries/filing.json",
                                             ".hydra/sources/files/filing/filing.pdf"]
    assert out["evidence"] == ".hydra/sources/files/filing/filing.pdf"
    assert (archive / out["evidence"]).read_bytes() == content
    recs = json.loads((archive / "collections/filing.json").read_text())          # plain hydra JSON
    assert all(set(r) == {"id", "schema", "properties"} for r in recs)
    assert all(r["properties"]["sources"] == ["filing-source"] for r in recs if r["schema"] != "Employment")
    doc = json.loads((archive / ".hydra/sources/entries/filing.json").read_text())[0]["properties"]
    assert doc["title"] == ["filing"] and doc["fileName"] == ["filing.pdf"] and doc["mimeType"] == ["application/pdf"]
    assert doc["contentHash"] == [hashlib.sha1(content).hexdigest()]          # the file itself
    assert "bodyText" not in doc                                            # the file is the evidence, not a copy of its text
    text_sha1 = hashlib.sha1(server.read_file(str(f), archive=str(archive))["text"].encode()).hexdigest()
    import pypdf
    assert doc["processingAgent"] == [f"hydra-mcp files.read (pypdf {pypdf.__version__})"]
    assert doc["notes"] == [f"sha1 of the text extracted from this file: {text_sha1}"]
    assert doc["publisher"] == ["Acme Times"]
    assert {c["value"] for c in claims(archive) if c["prop"] == "sources"} == {"case:filing-source"}
    assert server.read_file(str(f), archive=str(archive))["already_in_archive"]["source"] == "filing-source"
    with pytest.raises(ValueError, match="already saved as evidence 'filing-source'.*\"id\""):
        server.save(RECORDS[:1], {"file": str(f)}, archive=str(archive))


def test_records_that_break_the_rules_are_not_saved(archive, tmp_path):
    (tmp_path / "a.txt").write_text(TEXT)
    out = server.save([{"schema": "Persn", "properties": {"name": "X"}},
                       {"schema": "Membership", "properties": {"member": "X"}}], {"file": str(tmp_path / "a.txt")},
                      "a", archive=str(archive))
    assert out == {"saved": False, "errors": out["errors"], "warnings": out["warnings"]} and len(out["errors"]) == 2
    assert not (archive / "collections/a.json").exists() and not (archive / ".hydra/sources/files/a").exists()


def test_entities_already_in_the_archive_cite_the_new_evidence(archive, tmp_path):
    server.save(RECORDS, {"text": TEXT, "title": "first report"}, "first", archive=str(archive))
    (tmp_path / "minutes.txt").write_text("Lia Ortiz joined the Port Authority.")
    out = server.save([{"schema": "Person", "properties": {"name": "Lia Ortiz"}},
                       {"schema": "Membership", "properties": {"member": "Lia Ortiz", "organization": "Port Authority"}}],
                      {"file": str(tmp_path / "minutes.txt")}, "minutes", archive=str(archive))
    assert out["saved"] and out["existing_entities_cited"] == ["Port Authority"]
    assert not [i for i in out["issues"] if i["code"] == "prefixes.archive-collision"]        # no id taken twice
    linked = lambda name: [{l["caption"] for l in n["linked"]} for n in nodes_list(archive) if n["caption"] == name]
    assert linked("Port Authority") == [{"first report", "minutes"}] and linked("Dana Hale") == [{"first report"}]


def test_more_records_can_cite_evidence_already_saved(archive, tmp_path):
    (tmp_path / "a.txt").write_text(TEXT)
    server.save(RECORDS[:3], {"file": str(tmp_path / "a.txt")}, "a", archive=str(archive))
    out = server.save([RECORDS[4], RECORDS[2], RECORDS[0]], {"id": "case:a-source"}, "a-more", archive=str(archive))
    assert out["saved"] and out["files"] == ["collections/a-more.json"] and out["source_id"] == "a-source"
    with pytest.raises(ValueError, match="no source 'nope'"):
        server.save(RECORDS[:1], {"id": "nope"}, archive=str(archive))


def test_a_failed_build_leaves_no_copied_file_behind(archive, tmp_path, monkeypatch):
    (tmp_path / "a.txt").write_text(TEXT)
    monkeypatch.setattr(core.build, "build_archive", lambda root: (_ for _ in ()).throw(RuntimeError("boom")))
    with pytest.raises(ValueError, match="nothing was kept"):
        server.save(RECORDS[:1], {"file": str(tmp_path / "a.txt")}, "a", archive=str(archive))
    assert not (archive / "collections/a.json").exists() and not (archive / ".hydra/sources/files/a/a.txt").exists()
    assert not (archive / ".hydra/sources/entries/a.json").exists()


def test_a_file_can_be_read_before_any_archive_exists(tmp_path, monkeypatch):
    monkeypatch.delenv("HYDRA_ARCHIVE", raising=False)
    monkeypatch.chdir(tmp_path)
    (tmp_path / "a.txt").write_text(TEXT)
    assert server.read_file(str(tmp_path / "a.txt"))["text"] == TEXT
    with pytest.raises(ValueError, match="not a full path"):
        server.read_file("a.txt")
